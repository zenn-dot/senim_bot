"""Text extraction from PDF and DOCX. Never claims success when no text was read."""
from dataclasses import dataclass, field

MIN_TEXT_CHARS = 20  # fewer characters than this in the whole file = treated as no usable text


@dataclass
class ExtractResult:
    status: str                  # ok | empty | scanned | failed
    message: str = ""
    pages: list = field(default_factory=list)
    page_count: int | None = None   # None for DOCX (page numbers are not available)

    @property
    def text(self):
        return "\f".join(self.pages)


def extract_pdf(path) -> ExtractResult:
    pages = []
    try:
        try:
            import fitz  # PyMuPDF (preferred)
            with fitz.open(path) as doc:
                pages = [p.get_text() for p in doc]
        except ImportError:
            from pypdf import PdfReader  # fallback if PyMuPDF is not installed
            pages = [(p.extract_text() or "") for p in PdfReader(str(path)).pages]
    except Exception as e:
        return ExtractResult("failed", f"File is corrupted or not a readable PDF ({type(e).__name__}).")
    if not pages:
        return ExtractResult("empty", "PDF contains no pages.", [], 0)
    if sum(len(p.strip()) for p in pages) < MIN_TEXT_CHARS:
        return ExtractResult("scanned", "No text layer found. This looks like a scanned PDF; "
                             "OCR is not supported in this MVP, so it was NOT analyzed.", pages, len(pages))
    return ExtractResult("ok", "Text extracted.", pages, len(pages))


def extract_docx(path) -> ExtractResult:
    try:
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    except Exception as e:
        return ExtractResult("failed", f"File is corrupted or not a readable DOCX ({type(e).__name__}).")
    if len(text.strip()) < MIN_TEXT_CHARS:
        return ExtractResult("empty", "Document contains no usable text.", [text], None)
    return ExtractResult("ok", "Text extracted. DOCX has no fixed pages: page numbers unavailable.", [text], None)


def extract(path, filename) -> ExtractResult:
    name = filename.lower()
    if name.endswith(".pdf"):
        return extract_pdf(path)
    if name.endswith(".docx"):
        return extract_docx(path)
    return ExtractResult("failed", "Unsupported file type.")


def classify_document(filename: str, text: str) -> str:
    """Keyword heuristic used to pre-select a category; the user can always change it."""
    s = (filename + " " + text[:1500]).lower()
    rules = [("procedure", ["procedure", "instruction", "sop"]),
             ("work_permit", ["work permit", "permit to work"]),
             ("test_report", ["test report", "leak test", "inspection report", "test result"]),
             ("repair_report", ["repair completion", "completion report", "repair report"]),
             ("maintenance_record", ["maintenance record", "maintenance log"])]
    for cat, words in rules:
        if any(w in s for w in words):
            return cat
    return "other"
